"""GIC segment-wise premium workbook: fetch, parse, and insurer matching.

Reads the GIC ``.xlsx`` directly and returns exact **Gross Direct Premium
Income** figures (INR crore, usually provisional) for Indian non-life insurers.
Parsing fails loudly — a changed layout raises rather than returning plausible
wrong numbers.

No cloud-SDK dependency, so this unit-tests offline. The public tool entry point
(:func:`query_industry_premium`) and its hand-written schema live at the bottom.
"""

from __future__ import annotations

import json
import logging
import re
import time
from dataclasses import dataclass, field
from io import BytesIO

import requests

from shared import gic_reports

logger = logging.getLogger(__name__)

MEASURE = "Gross Direct Premium Income"
UNITS = "INR crore"

# key -> (exact sheet name, columns that MUST exist after the header is found)
SHEETS = {
    "segment": ("Segmentwise Report", ("Fire", "Motor Total", "Grand Total")),
    "health": ("Health Portfolio", ("Health-Retail", "Health-Group", "Grand Total")),
    "liability": ("Liability Portfolio", ("Product Liability", "Grand Total")),
    "miscellaneous": ("Miscellaneous portfolio", ("Crop Insurance", "Grand Total")),
}

# Genuine rebrands only. NEVER alias an acquired brand onto its acquirer — that
# silently misattributes figures.
_ALIASES = {
    "bajaj allianz": "Bajaj General Insurance Limited",
    "future generali": "Generali Central Insurance Company Limited",
    "kotak": "Zurich Kotak Mahindra General Insurance Co Ltd",
    "edelweiss": "Zuno General Insurance Co Ltd",
}

# Life-insurer detection. Must NOT fire on legitimate non-life names such as
# "The New India Assurance Co Ltd" or "Agriculture Insurance Co Of India Ltd".
_LIFE_RE = re.compile(
    r"\b(life|lic|max\s+life|annuity)\b|assurance\s+co\s+of\s+india",
    re.IGNORECASE,
)

_AGGREGATE_RE = re.compile(
    r"sub\s*total|subtotal|industry\s+total|%\s*growth|market\s+share|grand\s+total",
    re.IGNORECASE,
)

_PREVIOUS_YEAR_RE = re.compile(r"previous\s+year", re.IGNORECASE)

_WORKBOOK_TTL = 6 * 3600.0  # parsed workbook cached longer than the catalogue


class GicWorkbookError(RuntimeError):
    """Raised when the workbook cannot be fetched or its layout is unexpected."""


class LifeInsurerRequested(ValueError):
    """Raised when the query names a life insurer — this report is non-life only."""


@dataclass(frozen=True)
class InsurerRow:
    name: str
    values: dict
    previous_year: dict = field(default_factory=dict)


@dataclass(frozen=True)
class SheetData:
    sheet: str
    columns: list
    insurers: list
    aggregates: dict
    title: str


# ── fetch ─────────────────────────────────────────────────────────────

_SPREADSHEET_CTYPES = (
    "spreadsheetml",
    "application/vnd.ms-excel",
    "application/octet-stream",
    "application/zip",
)


def fetch_workbook(url: str, *, timeout: float = 60.0) -> bytes:
    """Fetch *url* and return workbook bytes, rejecting the SPA shell.

    A missing GIC file returns HTTP 200 with ``text/html`` (the app shell), so we
    check **both** the declared content type **and** the magic bytes: a real xlsx
    is a zip and starts with ``PK``; the shell starts with ``<``.
    """
    try:
        resp = requests.get(url, timeout=timeout)
        resp.raise_for_status()
    except requests.RequestException as exc:
        raise GicWorkbookError(f"Could not fetch workbook: {exc}") from exc

    ctype = resp.headers.get("Content-Type", "").lower()
    data = resp.content

    if "text/html" in ctype:
        raise GicWorkbookError(
            f"Expected a workbook but got HTML (Content-Type={ctype!r}) — the file "
            f"likely does not exist at {url}"
        )
    if not any(tok in ctype for tok in _SPREADSHEET_CTYPES) and ctype:
        # Unknown but non-HTML type: fall through to the magic-byte check.
        logger.warning("Unexpected workbook Content-Type %r for %s", ctype, url)
    if data[:2] != b"PK":
        raise GicWorkbookError(
            "Payload is not an .xlsx (missing 'PK' zip signature) — refusing to parse."
        )
    return data


def load_workbook(data: bytes):
    """Load workbook bytes with openpyxl (values only)."""
    import openpyxl  # imported lazily so discovery stays dependency-light

    try:
        return openpyxl.load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        raise GicWorkbookError(f"Could not open workbook: {exc}") from exc


# ── parse ─────────────────────────────────────────────────────────────

def _clean(text) -> str:
    """Collapse internal whitespace (fixes the double-spaced Marine headers)."""
    if text is None:
        return ""
    return re.sub(r"\s+", " ", str(text)).strip()


def _is_number(v) -> bool:
    if isinstance(v, bool):
        return False
    if isinstance(v, (int, float)):
        return True
    if isinstance(v, str):
        try:
            float(v.replace(",", ""))
            return True
        except ValueError:
            return False
    return False


def _to_number(v):
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, str):
        try:
            return float(v.replace(",", ""))
        except ValueError:
            return None
    return None


def find_header_row(rows, *, max_scan: int = 8) -> int:
    """Return the index of the header row (0-based) within the first *max_scan*.

    The header is the first row carrying **three or more non-empty text labels**
    to the right of column 0. Raises if none is found.
    """
    for i, row in enumerate(rows[:max_scan]):
        labels = 0
        for cell in row[1:]:
            text = _clean(cell)
            if text and not _is_number(cell):
                labels += 1
        if labels >= 3:
            return i
    raise GicWorkbookError("Could not locate a header row (no row with >=3 text labels).")


def parse_sheet(workbook, key: str) -> SheetData:
    """Parse the sheet identified by *key* into a :class:`SheetData`."""
    if key not in SHEETS:
        raise GicWorkbookError(f"Unknown sheet key: {key!r}")
    sheet_name, required = SHEETS[key]
    if sheet_name not in workbook.sheetnames:
        raise GicWorkbookError(
            f"Sheet {sheet_name!r} not found. Available: {list(workbook.sheetnames)}"
        )

    ws = workbook[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if not rows:
        raise GicWorkbookError(f"Sheet {sheet_name!r} is empty.")

    title = _clean(rows[0][0]) if rows[0] else ""
    header_idx = find_header_row(rows)
    header_cells = rows[header_idx]

    # Column name -> column index (skip col 0 which holds the insurer name).
    columns: list[str] = []
    col_index: dict[str, int] = {}
    for ci, cell in enumerate(header_cells):
        if ci == 0:
            continue
        name = _clean(cell)
        if name:
            columns.append(name)
            col_index.setdefault(name, ci)

    missing = [c for c in required if c not in col_index]
    if missing:
        raise GicWorkbookError(
            f"Sheet {sheet_name!r} missing required columns {missing}. "
            f"Found columns: {columns}"
        )

    insurers: list[InsurerRow] = []
    aggregates: dict[str, dict] = {}
    last_insurer_idx: int | None = None

    for row in rows[header_idx + 1:]:
        if not row:
            continue
        name = _clean(row[0])
        if not name:
            continue

        values = {}
        for col in columns:
            num = _to_number(row[col_index[col]]) if col_index[col] < len(row) else None
            if num is not None:
                values[col] = num

        # Section heading: a name with no numeric values. Resets PY pairing.
        if not values:
            last_insurer_idx = None
            continue

        # Aggregate row (sub total / industry total / % growth / market share…).
        if _AGGREGATE_RE.search(name):
            aggregates[name] = values
            last_insurer_idx = None
            continue

        # Unlabelled "Previous Year" row belongs to the insurer above it.
        if _PREVIOUS_YEAR_RE.search(name):
            if last_insurer_idx is not None:
                prev = insurers[last_insurer_idx]
                insurers[last_insurer_idx] = InsurerRow(
                    name=prev.name, values=prev.values, previous_year=values
                )
            continue

        insurers.append(InsurerRow(name=name, values=values))
        last_insurer_idx = len(insurers) - 1

    if len(insurers) < 5:
        raise GicWorkbookError(
            f"Sheet {sheet_name!r} parsed only {len(insurers)} insurers — layout looks wrong."
        )

    return SheetData(
        sheet=sheet_name, columns=columns, insurers=insurers,
        aggregates=aggregates, title=title,
    )


# ── insurer matching ──────────────────────────────────────────────────

def _tokens(text: str) -> list[str]:
    return [t for t in re.split(r"[^a-z0-9]+", text.lower()) if t]


def resolve_insurer(query: str, insurers: list) -> tuple[list, list]:
    """Resolve *query* to insurer rows.

    Returns ``(matches, candidates)``. Strict subset match — every query token
    must appear in the row name (no overlap scoring). Rebrands go through a small
    alias map. Life insurers raise :class:`LifeInsurerRequested`.
    """
    q = _clean(query)
    if _LIFE_RE.search(q):
        raise LifeInsurerRequested(
            f"{query!r} looks like a life insurer; this report covers non-life only."
        )

    # Alias rebrands to the registered name, then match that.
    match_target = q
    for alias, registered in _ALIASES.items():
        if alias in q.lower():
            match_target = registered
            break

    q_tokens = set(_tokens(match_target))
    if not q_tokens:
        return [], [row.name for row in insurers]

    matches = []
    for row in insurers:
        name_tokens = set(_tokens(row.name))
        if q_tokens <= name_tokens:  # subset
            matches.append(row)

    if matches:
        return matches, []
    return [], [row.name for row in insurers]


# ── tool: fetch + parse cache ─────────────────────────────────────────

_workbook_cache: dict[str, tuple[float, dict]] = {}
_workbook_last_good: dict[str, dict] = {}


def _get_sheet(period_key: str, url: str, sheet_key: str) -> SheetData:
    """Return parsed SheetData for (period, sheet_key), cached per period."""
    now = time.time()
    cached = _workbook_cache.get(period_key)
    if cached and (now - cached[0]) < _WORKBOOK_TTL and sheet_key in cached[1]:
        return cached[1][sheet_key]

    try:
        data = fetch_workbook(url)
        wb = load_workbook(data)
        parsed = parse_sheet(wb, sheet_key)
        store = cached[1] if (cached and (now - cached[0]) < _WORKBOOK_TTL) else {}
        store[sheet_key] = parsed
        _workbook_cache[period_key] = (now, store)
        _workbook_last_good.setdefault(period_key, {})[sheet_key] = parsed
        return parsed
    except Exception:
        good = _workbook_last_good.get(period_key, {})
        if sheet_key in good:
            logger.exception("Workbook fetch/parse failed for %s; using last-good", period_key)
            return good[sheet_key]
        raise


# ── hand-written schema + tool entry point ────────────────────────────

GIC_PREMIUM_SCHEMA = {
    "type": "object",
    "properties": {
        "insurers": {
            "type": "array", "items": {"type": "string"},
            "description": "Insurers by common or registered name. Omit for the whole industry.",
        },
        "segments": {
            "type": "array", "items": {"type": "string"},
            "description": "Columns to return, e.g. 'Motor Total', 'Fire', 'Grand Total'. Omit for all.",
        },
        "portfolio": {
            "type": "string",
            "enum": ["segment", "health", "liability", "miscellaneous"],
            "description": "Which sheet. 'segment' is the main one.",
        },
        "period": {
            "type": "string",
            "description": "e.g. 'July 2026'. OMIT for the latest published report, which is "
                           "the normal case. Never ask the user for a period, file name or URL.",
        },
        "include_previous_year": {"type": "boolean"},
        "top_n": {
            "type": "integer",
            "description": "Return only the largest N by the first requested segment, for league tables.",
        },
    },
    "required": [],
    "additionalProperties": False,
}

GIC_PREMIUM_DESCRIPTION = (
    "AUTHORITATIVE source for Indian NON-LIFE (general) insurance premium. Reads the "
    "General Insurance Council segment-wise workbook directly, giving exact published "
    "figures for every insurer. Use this, not web search, for any non-life premium, "
    "segment, ranking or market-share question. It finds the right report itself; omit "
    "period for the latest. Does NOT cover life insurers, and reports Gross Direct "
    "Premium, not Gross Written Premium."
)


def _select_columns(sheet: SheetData, segments) -> list:
    if not segments:
        return list(sheet.columns)
    chosen, unknown = [], []
    for s in segments:
        s_clean = _clean(s)
        if s_clean in sheet.columns:
            chosen.append(s_clean)
        else:
            unknown.append(s_clean)
    if not chosen:
        # Fall back to all columns rather than returning nothing.
        return list(sheet.columns)
    return chosen


def query_industry_premium(
    insurers=None, segments=None, portfolio="segment", period=None,
    include_previous_year=False, top_n=None, **_ignored,
) -> str:
    """Tool entry point. Returns a JSON string with provenance; fails closed."""
    try:
        portfolio = portfolio or "segment"
        if portfolio not in SHEETS:
            return json.dumps({"error": f"Unknown portfolio {portfolio!r}. "
                                        f"Choose one of {list(SHEETS)}."})

        resolution = gic_reports.resolve_report(gic_reports.SEGMENT, period)
        report = resolution.report
        sheet = _get_sheet(report.period, report.url, portfolio)

        columns = _select_columns(sheet, segments)

        payload = {
            "measure": MEASURE,
            "units": UNITS,
            "period": report.period_label,
            "provisional": True,
            "source_url": report.url,
            "source": "GIC segment-wise report",
            "sheet": sheet.sheet,
            "columns": columns,
            "rows": [],
        }
        if resolution.fell_back:
            payload["period_requested"] = resolution.requested
            payload["fell_back"] = True
            payload["note"] = (
                f"The requested period {resolution.requested!r} is not published; "
                f"quoting the latest ({report.period_label}). Say so to the user."
            )

        # Resolve insurers (or whole industry).
        selected_rows = []
        unmatched = []
        life_requested = []
        if insurers:
            for q in insurers:
                try:
                    matches, candidates = resolve_insurer(q, sheet.insurers)
                except LifeInsurerRequested:
                    life_requested.append(q)
                    continue
                if matches:
                    selected_rows.extend(matches)
                else:
                    unmatched.append({"requested": q, "candidates": candidates[:15]})
        else:
            selected_rows = list(sheet.insurers)

        # top_n ranking by the first requested column.
        if top_n and columns:
            key_col = columns[0]
            selected_rows = sorted(
                selected_rows, key=lambda r: r.values.get(key_col, float("-inf")),
                reverse=True,
            )[:top_n]

        for row in selected_rows:
            entry = {"insurer": row.name,
                     "values": {c: row.values.get(c) for c in columns if c in row.values}}
            if include_previous_year and row.previous_year:
                entry["previous_year"] = {c: row.previous_year.get(c)
                                          for c in columns if c in row.previous_year}
            payload["rows"].append(entry)

        # Whole-industry aggregates on request (no specific insurers).
        if not insurers and sheet.aggregates:
            payload["aggregates"] = {
                name: {c: vals.get(c) for c in columns if c in vals}
                for name, vals in sheet.aggregates.items()
            }

        if unmatched:
            payload["unmatched_insurers"] = unmatched
            payload["unmatched_note"] = (
                "Some names did not match. Ask the user which candidate was meant — do not guess."
            )
        if life_requested:
            payload["life_insurers_requested"] = life_requested
            payload["life_note"] = (
                "This report is non-life only. Tell the user it does not cover these life "
                "insurers; do NOT substitute a similarly named general insurer."
            )

        return json.dumps(payload, default=str)
    except Exception as exc:  # fail closed
        logger.exception("query_industry_premium failed")
        return json.dumps({"error": str(exc)})
